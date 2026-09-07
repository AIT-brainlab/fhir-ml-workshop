"""
Insert one page from a single-page Canva export into the finished deck.

    python3 slides/add_page.py extra.pptx --after 3

Copies the page's shapes, images and relationships verbatim, so the AIT design
survives intact, then places it at the requested position. Used to add the
"why not an LLM" page (template page 23) without re-exporting all 29 pages.

Needs python-pptx, which is not a workshop dependency: pip install python-pptx
"""

import argparse
import copy
import tempfile
from pathlib import Path

from pptx import Presentation

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL_ATTRS = (f"{R_NS}embed", f"{R_NS}link", f"{R_NS}id")

HERE = Path(__file__).resolve().parent
DECK = HERE / "from-fhir-to-a-table.pptx"


def copy_slide(src_slide, dst_prs):
    """Append a deep copy of src_slide to dst_prs, remapping its relationships."""
    layout = dst_prs.slides[0].slide_layout
    dst = dst_prs.slides.add_slide(layout)

    # the layout's own placeholders are not wanted; the source brings everything
    for shape in list(dst.shapes):
        shape._element.getparent().remove(shape._element)

    # Bring across every part the source slide points at, keeping a rId map.
    # Images are re-added from their bytes rather than by reference: the source
    # package numbers its media from image1.jpeg too, and reusing those part
    # objects puts two different parts under one name, which silently corrupts
    # the zip. get_or_add_image_part allocates an unused name and dedupes by
    # content hash.
    remap = {}
    for rid, rel in src_slide.part.rels.items():
        if rel.reltype.endswith("/slideLayout"):
            continue
        if rel.is_external:
            remap[rid] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype.endswith("/image"):
            blob = rel.target_part.blob
            ext = Path(str(rel.target_part.partname)).suffix or ".png"
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as fh:
                fh.write(blob)
                tmp = fh.name
            try:
                _, remap[rid] = dst.part.get_or_add_image_part(tmp)
            except Exception:
                # Vector art (EMF/WMF/SVG) that Pillow cannot open. Fall back to
                # referencing the part; these names rarely collide because the
                # bitmap media is what gets numbered from image1 upward.
                remap[rid] = dst.part.relate_to(rel.target_part, rel.reltype)
            finally:
                Path(tmp).unlink(missing_ok=True)
        else:
            remap[rid] = dst.part.relate_to(rel.target_part, rel.reltype)

    tree = dst.shapes._spTree
    for shape in src_slide.shapes:
        el = copy.deepcopy(shape._element)
        for node in el.iter():
            for attr in REL_ATTRS:
                old = node.get(attr)
                if old in remap:
                    node.set(attr, remap[old])
        tree.append(el)

    # background, if the source page carries one
    src_bg = src_slide._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
    )
    if src_bg is not None:
        # p:bg, not a:bg - a slide's background lives in the presentationml
        # namespace, and looking in the drawingml one silently finds nothing,
        # which drops the page onto white and hides any white text on it
        bg = src_bg.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
        if bg is not None:
            dst_cSld = dst._element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
            )
            dst_cSld.insert(0, copy.deepcopy(bg))

    return dst


def move_to(prs, from_index: int, to_index: int) -> None:
    """Move a slide (0-based) to another 0-based position."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    sid = ids[from_index]
    lst.remove(sid)
    lst.insert(to_index, sid)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("source", help=".pptx to take the layout from")
    ap.add_argument("--page", type=int, default=1,
                    help="1-based page inside the source file")
    ap.add_argument("--after", type=int, required=True,
                    help="1-based deck page the new page should follow")
    ap.add_argument("--deck", default=str(DECK))
    args = ap.parse_args()

    dst = Presentation(args.deck)
    src = Presentation(args.source)
    before = len(dst.slides._sldIdLst)

    copy_slide(src.slides[args.page - 1], dst)
    move_to(dst, before, args.after)       # appended last, now slot it in

    dst.save(args.deck)
    print(f"page inserted at position {args.after + 1}; deck now has "
          f"{len(Presentation(args.deck).slides._sldIdLst)} pages")


if __name__ == "__main__":
    main()
