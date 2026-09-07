"""
Refresh the finished deck's numbers from a real run.

    uv run python slides/collect_results.py
    uv run python slides/apply_results.py

Use this when the code or the data changed but the AIT template did not — it
rewrites the text of the existing deck in place, so you do not have to
re-export the template from Canva.

It reads exactly the same templates as build_from_template.py, so the two can
never disagree.
"""

import copy
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.text.text import _Paragraph

HERE = Path(__file__).resolve().parent
DECK = HERE / "from-fhir-to-a-table.pptx"

# import the sibling builder without needing it to be a package
_spec = importlib.util.spec_from_file_location(
    "builder", HERE / "build_from_template.py"
)
builder = importlib.util.module_from_spec(_spec)
sys.modules["builder"] = builder
_spec.loader.exec_module(builder)


def set_text(shape, text: str) -> None:
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        return
    lines = text.split("\n")
    p0.runs[0].text = lines[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for line in lines[1:]:
        new_p = copy.deepcopy(p0._p)
        p0._p.getparent().append(new_p)
        para = _Paragraph(new_p, tf)
        para.runs[0].text = line
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)


def main() -> None:
    if not DECK.exists():
        raise SystemExit(f"{DECK.name} not found — run build_from_template.py first.")

    facts = builder.load_results()
    prs = Presentation(DECK)

    # ORDER maps deck page -> the template page its text came from
    changed = 0
    for page, origin in enumerate(builder.ORDER, start=1):
        mapping = builder.TEXT.get(origin)
        if not mapping:
            continue
        slide = prs.slides[page - 1]
        by_name = {sh.name: sh for sh in slide.shapes}
        for name, template in mapping.items():
            shape = by_name.get(name)
            if shape is None:
                continue
            want = template.format(**facts)
            if shape.text_frame.text != want:
                set_text(shape, want)
                changed += 1
        if origin in builder.NOTES:
            slide.notes_slide.notes_text_frame.text = builder.NOTES[origin]

    prs.save(DECK)
    print(f"{changed} text box(es) refreshed from slides/results.json")


if __name__ == "__main__":
    main()
